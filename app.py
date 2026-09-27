import os
import json
import base64
import sqlite3
import re
import io
from datetime import datetime
from functools import wraps

from flask import (
    Flask,
    request,
    redirect,
    url_for,
    session,
    flash,
    render_template_string,
    send_file,
    send_from_directory,
)

from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
import fitz
from openai import OpenAI


# ============================================================
# CONFIGURATION
# ============================================================

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change-this-secret-key-in-render")

DATABASE = os.getenv("DATABASE_PATH", "platform.db")
UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "uploads")
AI_MODEL = os.getenv("AI_MODEL", "gpt-5.6-luna")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin123")

MAX_FILE_SIZE_MB = 20
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE_MB * 1024 * 1024
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
client = OpenAI(api_key=OPENAI_API_KEY) if OPENAI_API_KEY else None


# ============================================================
# DATABASE
# ============================================================

def get_db():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    return connection


def init_database():
    db = get_db()

    db.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            full_name TEXT NOT NULL,
            stage TEXT NOT NULL,
            study_type TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    """)

    db.execute("""
        CREATE TABLE IF NOT EXISTS settings (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            site_name TEXT NOT NULL,
            logo_url TEXT DEFAULT '',
            primary_color TEXT DEFAULT '#0d6efd',
            secondary_color TEXT DEFAULT '#07111f',
            university_notice TEXT DEFAULT 'جامعة الأنبار'
        )
    """)

    db.execute("""
        CREATE TABLE IF NOT EXISTS announcements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    """)

    db.execute("""
        CREATE TABLE IF NOT EXISTS generated_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER,
            item_type TEXT NOT NULL,
            title TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    """)

    existing = db.execute(
        "SELECT id FROM settings WHERE id = 1"
    ).fetchone()

    if not existing:
        db.execute("""
            INSERT INTO settings
            (id, site_name, logo_url, primary_color, secondary_color, university_notice)
            VALUES
            (1, 'منصة الأنبار التعليمية الذكية', '', '#0d6efd', '#07111f', 'جامعة الأنبار')
        """)

    db.commit()
    db.close()


init_database()


# ============================================================
# HELPERS
# ============================================================

def get_settings():
    db = get_db()
    settings = db.execute(
        "SELECT * FROM settings WHERE id = 1"
    ).fetchone()
    db.close()
    return settings


def login_required(function):
    @wraps(function)
    def wrapper(*args, **kwargs):
        if "student_id" not in session:
            return redirect(url_for("attendance"))
        return function(*args, **kwargs)
    return wrapper


def admin_required(function):
    @wraps(function)
    def wrapper(*args, **kwargs):
        if not session.get("admin_logged_in"):
            return redirect(url_for("admin_login"))
        return function(*args, **kwargs)
    return wrapper


def allowed_file(filename):
    extensions = {
        "pdf", "png", "jpg", "jpeg", "webp",
        "txt", "md", "csv", "docx", "pptx", "xlsx"
    }
    return "." in filename and filename.rsplit(".", 1)[1].lower() in extensions


def extract_pdf_text(file_path):
    try:
        reader = PdfReader(file_path)
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages)[:70000]
    except Exception as error:
        return f"PDF extraction error: {error}"


def extract_docx_text(file_path):
    document = Document(file_path)
    parts = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)[:70000]


def extract_pptx_text(file_path):
    presentation = Presentation(file_path)
    parts = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        slide_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if slide_parts:
            parts.append(f"[Slide {slide_number}]\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)[:70000]


def extract_xlsx_text(file_path):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                parts.append(" | ".join(values))
    workbook.close()
    return "\n".join(parts)[:70000]


def extract_text_file(file_path):
    with open(file_path, "r", encoding="utf-8", errors="replace") as file:
        return file.read()[:70000]


def extract_json(text):
    text = text.strip()
    text = re.sub(r"^```json\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^```\s*", "", text)
    text = re.sub(r"\s*```$", "", text)

    match = re.search(r"\{.*\}", text, flags=re.DOTALL)
    return json.loads(match.group(0) if match else text)


def ai_text(prompt):
    if not client:
        raise RuntimeError("OPENAI_API_KEY is not configured.")

    response = client.responses.create(
        model=AI_MODEL,
        input=prompt
    )
    return response.output_text


def ai_image(image_path, prompt):
    if not client:
        raise RuntimeError("OPENAI_API_KEY is not configured.")

    extension = os.path.splitext(image_path)[1].lower()
    mime = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp"
    }.get(extension, "image/jpeg")

    with open(image_path, "rb") as image_file:
        encoded = base64.b64encode(image_file.read()).decode("utf-8")

    image_url = f"data:{mime};base64,{encoded}"

    response = client.responses.create(
        model=AI_MODEL,
        input=[{
            "role": "user",
            "content": [
                {"type": "input_text", "text": prompt},
                {"type": "input_image", "image_url": image_url}
            ]
        }]
    )
    return response.output_text


def process_uploaded_file(uploaded_file, prompt):
    if not uploaded_file or not uploaded_file.filename:
        raise ValueError("لم يتم اختيار ملف أو صورة.")

    filename = secure_filename(uploaded_file.filename)
    if not filename:
        raise ValueError("اسم الملف غير صالح.")

    if not allowed_file(filename):
        raise ValueError(
            "الملفات المدعومة: PDF, PNG, JPG, JPEG, WEBP, DOCX, PPTX, XLSX, TXT, MD, CSV."
        )

    saved_name = datetime.now().strftime("%Y%m%d%H%M%S%f") + "_" + filename
    file_path = os.path.join(UPLOAD_FOLDER, saved_name)
    uploaded_file.save(file_path)
    extension = os.path.splitext(filename)[1].lower()

    # Images are sent directly to the vision-capable AI model.
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        return ai_image(file_path, prompt)

    # Office/text files are extracted on the server and then analyzed by AI.
    if extension == ".pdf":
        extracted = extract_pdf_text(file_path)
        if extracted.strip() and not extracted.startswith("PDF extraction error:"):
            return ai_text(prompt + "\n\nSOURCE MATERIAL:\n" + extracted)

        # Scanned/image-only PDFs: render pages and send them to the vision model.
        pdf = fitz.open(file_path)
        page_results = []
        page_count = len(pdf)
        max_pages = min(page_count, 20)
        try:
            for page_index in range(max_pages):
                page = pdf.load_page(page_index)
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                image_path = os.path.join(
                    UPLOAD_FOLDER,
                    f"pdf_page_{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{page_index}.png"
                )
                pix.save(image_path)
                try:
                    page_prompt = prompt + f"\n\nهذه الصفحة رقم {page_index + 1} من ملف PDF مصوّر. اقرأ النص الظاهر فيها بدقة ثم نفّذ المطلوب."
                    page_results.append(ai_image(image_path, page_prompt))
                finally:
                    try:
                        os.remove(image_path)
                    except OSError:
                        pass
        finally:
            pdf.close()

        if not page_results:
            raise ValueError("تعذر قراءة صفحات الـPDF المصوّر.")

        notice = "\n\n[تمت معالجة أول 20 صفحة من PDF المصوّر]\n" if page_count > 20 else "\n\n"
        return notice.join(page_results)

    if extension == ".docx":
        extracted = extract_docx_text(file_path)
    elif extension == ".pptx":
        extracted = extract_pptx_text(file_path)
    elif extension == ".xlsx":
        extracted = extract_xlsx_text(file_path)
    else:
        extracted = extract_text_file(file_path)

    if not extracted.strip():
        raise ValueError("الملف فارغ أو لم يتم العثور على نص قابل للقراءة.")

    return ai_text(prompt + "\n\nSOURCE MATERIAL:\n" + extracted)


def save_generated_item(student_id, item_type, title, content):
    db = get_db()
    cursor = db.execute("""
        INSERT INTO generated_items
        (student_id, item_type, title, content, created_at)
        VALUES (?, ?, ?, ?, ?)
    """, (
        student_id,
        item_type,
        title,
        content,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ))
    item_id = cursor.lastrowid
    db.commit()
    db.close()
    return item_id


def bilingual_rules():
    return """
IMPORTANT BILINGUAL RULES:
1. The student interface is Arabic.
2. Do not leave important English words or technical terms without Arabic meaning.
3. When an English term appears, keep the English term and immediately provide its Arabic meaning in parentheses.
   Example: Cell (الخلية).
4. For full English sentences, provide the Arabic translation directly below or after the English sentence.
5. Preserve scientific terminology, formulas, numbers and units.
6. Do not invent information.
"""


def export_text(title, content):
    text = f"{title}\n{'=' * 60}\n\n{content}"
    return send_file(
        io.BytesIO(text.encode("utf-8")),
        mimetype="text/plain; charset=utf-8",
        as_attachment=True,
        download_name="anbar-ai-result.txt"
    )


# ============================================================
# SHARED HTML
# ============================================================

UPLOAD_PROCESS_SCRIPT = """
<script>
document.querySelectorAll('form[enctype="multipart/form-data"]').forEach(function(form) {
    form.addEventListener('submit', function() {
        const button = form.querySelector('button[type="submit"]');
        const input = form.querySelector('input[type="file"]');
        if (input && !input.files.length) {
            alert('اختر ملفاً أو صورة أولاً.');
            return;
        }
        if (button) {
            button.disabled = true;
            button.innerHTML = '⏳ جاري رفع الملف وتحليله بالذكاء الاصطناعي...';
        }
        const status = document.createElement('div');
        status.className = 'alert';
        status.innerHTML = '🤖 تم استلام الملف. انتظر حتى يكتمل التحليل ثم ستظهر النتيجة هنا.';
        form.parentNode.insertBefore(status, form.nextSibling);
    });
});
</script>
"""

BASE_STYLE = """
<style>
:root {
    --primary: {{ settings.primary_color }};
    --secondary: {{ settings.secondary_color }};
    --background: #f4f7fb;
    --card: #ffffff;
    --text: #152238;
    --muted: #667085;
    --border: #e5e7eb;
}
* { box-sizing: border-box; }
body {
    margin: 0;
    font-family: Tahoma, Arial, sans-serif;
    background: linear-gradient(135deg, #f8fbff, #eef4fb);
    color: var(--text);
    direction: rtl;
}
a { text-decoration: none; color: inherit; }
.container { width: min(1150px, 94%); margin: auto; }
.navbar {
    background: var(--secondary);
    color: white;
    padding: 15px 0;
    position: sticky;
    top: 0;
    z-index: 100;
    box-shadow: 0 4px 20px rgba(0,0,0,.12);
}
.nav-inner {
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 15px;
}
.brand { display: flex; align-items: center; gap: 10px; font-weight: bold; }
.brand img {
    width: 42px; height: 42px; border-radius: 50%;
    object-fit: cover; background: white;
}
.nav-links { display: flex; gap: 8px; flex-wrap: wrap; }
.nav-links a { padding: 8px 12px; border-radius: 10px; font-size: 14px; }
.nav-links a:hover { background: rgba(255,255,255,.12); }
.page { padding: 30px 0 60px; }
.hero {
    background: linear-gradient(135deg, var(--secondary), #15365c);
    color: white; border-radius: 25px; padding: 35px;
    margin-bottom: 25px; box-shadow: 0 15px 40px rgba(0,0,0,.12);
}
.hero h1 { margin: 0 0 10px; font-size: clamp(25px, 5vw, 42px); }
.hero p { opacity: .9; line-height: 1.9; }
.grid {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
    gap: 18px;
}
.card {
    background: var(--card); border: 1px solid var(--border);
    border-radius: 20px; padding: 22px;
    box-shadow: 0 8px 30px rgba(15,23,42,.06);
}
.feature-card { transition: .25s; }
.feature-card:hover { transform: translateY(-5px); box-shadow: 0 15px 35px rgba(15,23,42,.12); }
.icon {
    width: 55px; height: 55px; border-radius: 16px;
    display: grid; place-items: center; background: #edf5ff;
    font-size: 25px; margin-bottom: 15px;
}
.btn {
    display: inline-flex; align-items: center; justify-content: center;
    border: 0; border-radius: 12px; padding: 12px 18px;
    cursor: pointer; background: var(--primary); color: white;
    font-weight: bold; margin: 6px 4px 0 0;
}
.btn:hover { opacity: .9; }
.btn.secondary { background: #eef2f7; color: #1f2937; }
.btn.dark { background: #172033; color: white; }
.form-card { max-width: 700px; margin: 30px auto; }
.form-group { margin-bottom: 18px; }
label { display: block; margin-bottom: 7px; font-weight: bold; }
input, select, textarea {
    width: 100%; border: 1px solid #d8dee8; border-radius: 12px;
    padding: 13px; font-size: 15px; outline: none; background: white;
}
input:focus, select:focus, textarea:focus {
    border-color: var(--primary); box-shadow: 0 0 0 3px rgba(13,110,253,.08);
}
textarea { min-height: 180px; resize: vertical; }
.upload-box {
    border: 2px dashed #b8c4d6; padding: 30px; text-align: center;
    border-radius: 18px; background: #fafcff; margin-bottom: 20px;
}
.result {
    white-space: pre-wrap; line-height: 2; background: #fbfdff;
    border: 1px solid #e4eaf2; border-radius: 16px;
    padding: 20px; margin-top: 20px; direction: rtl;
}
.alert {
    padding: 14px 17px; border-radius: 13px; margin-bottom: 18px;
    background: #eef6ff; border: 1px solid #cfe3ff;
}
.alert.error { background: #fff1f2; border-color: #fecdd3; }
.loading-screen {
    position: fixed; inset: 0;
    background: radial-gradient(circle at top, #17345b, #040a12);
    z-index: 9999; display: flex; align-items: center;
    justify-content: center; color: white;
}
.loading-content { text-align: center; width: 90%; max-width: 450px; }
.loading-logo {
    width: 100px; height: 100px; object-fit: cover; border-radius: 50%;
    background: white; padding: 5px; margin-bottom: 20px;
}
.loader {
    height: 7px; width: 100%; background: rgba(255,255,255,.15);
    border-radius: 20px; overflow: hidden; margin-top: 25px;
}
.loader span {
    display: block; height: 100%; width: 0; background: var(--primary);
    animation: loading 5s linear forwards;
}
@keyframes loading { to { width: 100%; } }
.stat { text-align: center; }
.stat strong { display: block; font-size: 30px; color: var(--primary); }
.admin-table { width: 100%; border-collapse: collapse; min-width: 760px; }
.admin-table th, .admin-table td {
    border-bottom: 1px solid #e5e7eb; padding: 12px; text-align: right;
}
.exam-question {
    background: white; border: 1px solid var(--border);
    border-radius: 18px; padding: 20px; margin-bottom: 15px;
}
.option {
    display: block; margin: 8px 0; padding: 12px;
    background: #f8fafc; border-radius: 10px; cursor: pointer;
}
.mind-map { display: flex; flex-direction: column; align-items: center; gap: 15px; }
.mind-node {
    padding: 15px 25px; background: white; border: 2px solid var(--primary);
    border-radius: 15px; box-shadow: 0 5px 15px rgba(0,0,0,.08);
    text-align: center; max-width: 95%;
}
.summary-card {
    background: white; border-radius: 20px; padding: 25px;
    border: 1px solid #e1e7ef; box-shadow: 0 10px 30px rgba(15,23,42,.08);
}
.notice {
    padding: 12px 15px; border-radius: 12px;
    background: #fff8e7; border: 1px solid #f4d58d; margin: 12px 0;
}
.footer {
    background: var(--secondary); color: white; padding: 25px 0;
    text-align: center; margin-top: 50px;
}
.kpi {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(160px, 1fr));
    gap: 12px; margin-bottom: 20px;
}
.kpi .card { padding: 18px; }
.small { color: var(--muted); font-size: 13px; }
@media(max-width:700px) {
    .nav-inner { flex-direction: column; }
    .nav-links { justify-content: center; }
    .hero { padding: 25px; }
}
</style>
"""


def result_actions(item_id, element_id, title="نتيجة الذكاء الاصطناعي"):
    if not item_id:
        return ""
    return f"""
<div class="notice">
    يمكنك حفظ النتيجة كملف أو صورة، أو مشاركتها من هاتفك.
</div>
<a class="btn secondary" href="{url_for('download_item', item_id=item_id)}">
    📄 تحميل كملف
</a>
<button class="btn" type="button" onclick="downloadCard('{element_id}')">
    🖼️ تحميل كصورة
</button>
<button class="btn dark" type="button" onclick="shareCard('{element_id}', {json.dumps(title, ensure_ascii=False)})">
    📤 مشاركة
</button>
<script>
async function downloadCard(id) {{
    if (!window.html2canvas) {{
        alert("انتظر تحميل مكتبة الصور ثم حاول مرة أخرى.");
        return;
    }}
    const canvas = await html2canvas(document.getElementById(id), {{
        backgroundColor: "#ffffff",
        scale: 2
    }});
    const link = document.createElement("a");
    link.download = "anbar-ai-result.png";
    link.href = canvas.toDataURL("image/png");
    link.click();
}}
async function shareCard(id, title) {{
    const element = document.getElementById(id);
    const text = element.innerText || "";
    if (navigator.share) {{
        try {{
            await navigator.share({{title: title, text: text}});
            return;
        }} catch (e) {{}}
    }}
    try {{
        await navigator.clipboard.writeText(text);
        alert("تم نسخ النتيجة. يمكنك إرسالها للطالب.");
    }} catch (e) {{
        alert("يمكنك استخدام زر تحميل الصورة أو الملف.");
    }}
}}
</script>
<script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script>
"""


def render_page(title, body, **context):
    settings = get_settings()

    template = f"""
<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title} | {{{{ settings.site_name }}}}</title>
<meta name="description" content="منصة تعليمية ذكية لطلاب جامعة الأنبار">
{BASE_STYLE}
</head>
<body>
<nav class="navbar">
<div class="container nav-inner">
<a class="brand" href="{{{{ url_for('home') }}}}">
{{% if settings.logo_url %}}
<img src="{{{{ settings.logo_url }}}}" alt="logo">
{{% else %}}
<div style="width:42px;height:42px;border-radius:50%;background:#fff;color:#123;display:grid;place-items:center;font-weight:bold;">AI</div>
{{% endif %}}
<span>{{{{ settings.site_name }}}}</span>
</a>
<div class="nav-links">
<a href="{{{{ url_for('home') }}}}">الرئيسية</a>
{{% if session.get('student_id') %}}
<a href="{{{{ url_for('translation') }}}}">الترجمة</a>
<a href="{{{{ url_for('summary') }}}}">الملخصات</a>
<a href="{{{{ url_for('mindmap') }}}}">المخططات</a>
<a href="{{{{ url_for('exam') }}}}">الاختبارات</a>
<a href="{{{{ url_for('announcements') }}}}">الإعلانات</a>
<a href="{{{{ url_for('logout') }}}}">خروج</a>
{{% endif %}}
</div>
</div>
</nav>
<main class="page">
<div class="container">
{{% with messages = get_flashed_messages(with_categories=true) %}}
{{% for category, message in messages %}}
<div class="alert {{{{ 'error' if category == 'error' else '' }}}}">{{{{ message }}}}</div>
{{% endfor %}}
{{% endwith %}}
{body}
</div>
</main>
{UPLOAD_PROCESS_SCRIPT}
<footer class="footer">
<div class="container">
{{{{ settings.university_notice }}}}<br>
<span>منصة تعليمية ذكية</span>
</div>
</footer>
</body>
</html>
"""
    return render_template_string(template, settings=settings, **context)


# ============================================================
# LOADING + ATTENDANCE
# ============================================================

@app.route("/")
def root():
    return redirect(url_for("loading"))


@app.route("/loading")
def loading():
    settings = get_settings()
    logo = (
        f"<img class='loading-logo' src='{settings['logo_url']}' alt='logo'>"
        if settings["logo_url"]
        else "<div class='loading-logo' style='display:grid;place-items:center;color:#123;font-size:28px;font-weight:bold;'>AI</div>"
    )
    html = f"""
<div class="loading-screen">
<div class="loading-content">
{logo}
<h1>{settings['site_name']}</h1>
<p>{settings['university_notice']}</p>
<p>جاري تجهيز المنصة التعليمية...</p>
<div class="loader"><span></span></div>
<p>يرجى الانتظار 5 ثوانٍ</p>
</div>
</div>
<script>
setTimeout(function() {{
    window.location.href = "{url_for('attendance')}";
}}, 5000);
</script>
"""
    return render_template_string(
        f"""<!DOCTYPE html><html lang="ar" dir="rtl"><head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{settings['site_name']}</title>{BASE_STYLE}</head><body>{html}</body></html>""",
        settings=settings
    )


@app.route("/attendance", methods=["GET", "POST"])
def attendance():
    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        stage = request.form.get("stage", "").strip()
        study_type = request.form.get("study_type", "").strip()

        if len(full_name.split()) < 3:
            flash("يرجى إدخال الاسم الثلاثي الكامل.", "error")
            return redirect(url_for("attendance"))

        if not stage:
            flash("يرجى اختيار المرحلة.", "error")
            return redirect(url_for("attendance"))

        if study_type not in ["morning", "evening"]:
            flash("يرجى اختيار نوع الدراسة.", "error")
            return redirect(url_for("attendance"))

        db = get_db()
        cursor = db.execute("""
            INSERT INTO students (full_name, stage, study_type, created_at)
            VALUES (?, ?, ?, ?)
        """, (
            full_name, stage, study_type,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ))
        student_id = cursor.lastrowid
        db.commit()
        db.close()

        session["student_id"] = student_id
        session["student_name"] = full_name
        session["student_stage"] = stage
        session["student_study_type"] = study_type

        return redirect(url_for("home"))

    body = """
<div class="card form-card">
<h2>تسجيل الحضور الإجباري</h2>
<p>قبل الدخول إلى المنصة يجب تسجيل الحضور.</p>
<form method="POST">
<div class="form-group">
<label>اسم الطالب الثلاثي</label>
<input type="text" name="full_name" placeholder="مثال: محمد علي حسن" required>
</div>
<div class="form-group">
<label>المرحلة الدراسية</label>
<select name="stage" required>
<option value="">اختر المرحلة</option>
<option>المرحلة الأولى</option><option>المرحلة الثانية</option>
<option>المرحلة الثالثة</option><option>المرحلة الرابعة</option>
<option>المرحلة الخامسة</option><option>المرحلة السادسة</option>
<option>دراسات عليا</option>
</select>
</div>
<div class="form-group">
<label>نوع الدراسة</label>
<select name="study_type" required>
<option value="">اختر نوع الدراسة</option>
<option value="morning">☀️ صباحي</option>
<option value="evening">🌙 مسائي</option>
</select>
</div>
<button class="btn" type="submit">تسجيل الحضور والدخول</button>
</form>
</div>
"""
    return render_page("تسجيل الحضور", body)


# ============================================================
# HOME
# ============================================================

@app.route("/home")
@login_required
def home():
    body = """
<div class="hero">
<h1>منصة التعليم الذكي</h1>
<p>
أهلاً بك <strong>{{ session.get('student_name') }}</strong><br>
منصة تعليمية تساعدك على ترجمة المحاضرات، تلخيصها،
تحويلها إلى مخططات وإنشاء اختبارات ثنائية اللغة.
</p>
</div>
<div class="grid">
<a class="card feature-card" href="{{ url_for('translation') }}"><div class="icon">🌍</div><h2>الترجمة الذكية</h2><p>ترجمة المحتوى مع إبقاء المصطلح الإنجليزي وإضافة معناه العربي.</p></a>
<a class="card feature-card" href="{{ url_for('summary') }}"><div class="icon">📚</div><h2>الملخصات</h2><p>ملخص عربي واضح مع ترجمة المصطلحات الإنجليزية المهمة.</p></a>
<a class="card feature-card" href="{{ url_for('mindmap') }}"><div class="icon">🧠</div><h2>المخططات</h2><p>مخططات ذهنية مرتبة بالعربي مع المصطلحات الإنجليزية ومعانيها.</p></a>
<a class="card feature-card" href="{{ url_for('exam') }}"><div class="icon">📝</div><h2>الاختبارات</h2><p>10 أسئلة جامعية، السؤال بالعربي والإنكليزي مع التصحيح.</p></a>
<a class="card feature-card" href="{{ url_for('announcements') }}"><div class="icon">📢</div><h2>إعلانات الجامعة</h2><p>آخر الإعلانات والتنبيهات الخاصة بالمنصة.</p></a>
<div class="card"><div class="icon">🤖</div><h2>AI</h2><p>معالجة PDF والصور مباشرة من الموقع.</p></div>
</div>
"""
    return render_page("الرئيسية", body)


# ============================================================
# TRANSLATION
# ============================================================

@app.route("/translation", methods=["GET", "POST"])
@login_required
def translation():
    result = None
    item_id = None

    if request.method == "POST":
        target_language = request.form.get("target_language", "Arabic")
        uploaded = request.files.get("file")

        try:
            prompt = f"""
You are a professional university translation and OCR assistant.
Target language: {target_language}

{bilingual_rules()}

The student wants the result to be immediately useful for study.
If the source contains English:
- translate every meaningful English word, phrase, sentence and technical term into Arabic;
- keep the original English term next to its Arabic meaning;
- organize the result by headings and paragraphs.

If the source is an image, read visible text accurately before translating it.
If the source is Arabic and contains English terms, translate those English terms into Arabic too.
Return the complete translated and organized result only.
"""
            result = process_uploaded_file(uploaded, prompt)
            item_id = save_generated_item(
                session["student_id"], "translation",
                "AI Translation", result
            )
        except Exception as error:
            result = "حدث خطأ أثناء الترجمة:\n" + str(error)

    body = """
<div class="hero">
<h1>🌍 الترجمة الذكية</h1>
<p>ارفع PDF أو صورة، وسيتم استخراج المحتوى وترجمته مع إظهار الإنجليزية والعربية معاً.</p>
</div>
<div class="card">
<form method="POST" enctype="multipart/form-data">
<div class="upload-box">
<h3>📎 رفع المحاضرة</h3>
<input type="file" name="file" accept=".pdf,.png,.jpg,.jpeg,.webp,.docx,.pptx,.xlsx,.txt,.md,.csv" required>
<p class="small">PDF أو صورة أو Word أو PowerPoint أو Excel أو TXT/CSV/MD — الحد الأقصى 20MB</p>
</div>
<div class="form-group">
<label>اللغة المطلوبة</label>
<select name="target_language">
<option value="Arabic">العربية</option>
<option value="English">English</option>
<option value="French">Français</option>
<option value="Turkish">Türkçe</option>
<option value="German">Deutsch</option>
</select>
</div>
<button class="btn" type="submit">🚀 ترجمة الملف الآن</button>
</form>
{% if result %}
<div class="summary-card" id="translationCard" style="margin-top:25px;">
<h2>🌍 النتيجة</h2>
<div class="result">{{ result }}</div>
</div>
{{ actions|safe }}
{% endif %}
</div>
"""
    actions = result_actions(item_id, "translationCard", "ترجمة المحاضرة")
    return render_page("الترجمة", body, result=result, actions=actions)


# ============================================================
# SUMMARY
# ============================================================

@app.route("/summary", methods=["GET", "POST"])
@login_required
def summary():
    result = None
    item_id = None

    if request.method == "POST":
        uploaded = request.files.get("file")
        try:
            prompt = f"""
You are an expert university academic summarization assistant.
Analyze the complete lecture.

{bilingual_rules()}

Create a clear Arabic study summary.
Requirements:
1. Major topics and headings.
2. Important concepts and definitions.
3. Difficult concepts explained simply.
4. Bullet points.
5. Important comparisons and formulas.
6. For every important English word/term, write English + Arabic meaning.
7. If a full English sentence is important, show the English sentence and its Arabic translation.
8. Finish with: Quick Review (مراجعة سريعة), Important Terms (المصطلحات المهمة), Important Points (النقاط المهمة).
Return clean study-ready text.
"""
            result = process_uploaded_file(uploaded, prompt)
            item_id = save_generated_item(
                session["student_id"], "summary",
                "AI Summary", result
            )
        except Exception as error:
            result = "حدث خطأ أثناء إنشاء الملخص:\n" + str(error)

    body = """
<div class="hero">
<h1>📚 الملخصات الذكية</h1>
<p>ارفع المحاضرة وسيتم إنشاء ملخص عربي واضح مع ترجمة المصطلحات الإنجليزية.</p>
</div>
<div class="card">
<form method="POST" enctype="multipart/form-data">
<div class="upload-box">
<h3>📎 رفع المحاضرة</h3>
<input type="file" name="file" accept=".pdf,.png,.jpg,.jpeg,.webp,.docx,.pptx,.xlsx,.txt,.md,.csv" required>
</div>
<button class="btn" type="submit">✨ إنشاء الملخص</button>
</form>
{% if result %}
<div class="summary-card" id="summaryCard" style="margin-top:25px;">
<h2>📚 الملخص</h2>
<div class="result">{{ result }}</div>
</div>
{{ actions|safe }}
{% endif %}
</div>
"""
    actions = result_actions(item_id, "summaryCard", "ملخص المحاضرة")
    return render_page("الملخصات", body, result=result, actions=actions)


# ============================================================
# MIND MAP
# ============================================================

@app.route("/mindmap", methods=["GET", "POST"])
@login_required
def mindmap():
    result = None
    item_id = None

    if request.method == "POST":
        uploaded = request.files.get("file")
        try:
            prompt = f"""
You are a university academic mind-map assistant.

{bilingual_rules()}

Analyze the entire lecture and return ONLY valid JSON:
{{
  "title": "Arabic title",
  "branches": [
    {{
      "title": "Arabic branch title + English term when useful",
      "points": [
        "Arabic point with English term and Arabic meaning when useful"
      ]
    }}
  ]
}}
Requirements:
- Cover the important parts of the lecture.
- Keep the map concise and educational.
- Every English term must have its Arabic meaning.
- Do not invent information.
"""
            raw = process_uploaded_file(uploaded, prompt)
            result = extract_json(raw)
            item_id = save_generated_item(
                session["student_id"], "mindmap",
                "AI Mind Map", json.dumps(result, ensure_ascii=False)
            )
        except Exception as error:
            result = {
                "title": "خطأ في المخطط",
                "branches": [{"title": "Error", "points": [str(error)]}]
            }

    body = """
<div class="hero">
<h1>🧠 المخططات الذكية</h1>
<p>مخطط ذهني منظم مع العربية وترجمة المصطلحات الإنجليزية.</p>
</div>
<div class="card">
<form method="POST" enctype="multipart/form-data">
<div class="upload-box">
<input type="file" name="file" accept=".pdf,.png,.jpg,.jpeg,.webp,.docx,.pptx,.xlsx,.txt,.md,.csv" required>
</div>
<button class="btn" type="submit">🧠 إنشاء المخطط</button>
</form>
{% if result %}
<div class="summary-card" id="mindMapCard" style="margin-top:25px;">
<h2>🧠 {{ result.title }}</h2>
<div class="mind-map">
{% for branch in result.branches %}
<div class="mind-node">
<strong>{{ branch.title }}</strong>
{% for point in branch.points %}
<div style="margin-top:8px;padding:8px;background:#f5f8fc;border-radius:8px;">{{ point }}</div>
{% endfor %}
</div>
{% endfor %}
</div>
</div>
{{ actions|safe }}
{% endif %}
</div>
"""
    actions = result_actions(item_id, "mindMapCard", "المخطط الذهني")
    return render_page("المخططات", body, result=result, actions=actions)


# ============================================================
# EXAM
# ============================================================

@app.route("/exam", methods=["GET", "POST"])
@login_required
def exam():
    exam_data = None
    item_id = None

    if request.method == "POST":
        uploaded = request.files.get("file")
        try:
            prompt = f"""
You are an experienced university examination designer.

{bilingual_rules()}

Analyze the entire uploaded academic material.
Create exactly 10 methodology-based university questions covering different sections.

Return ONLY valid JSON:
{{
  "title": "Arabic exam title",
  "questions": [
    {{
      "number": 1,
      "type": "mcq",
      "question_ar": "السؤال بالعربية",
      "question_en": "Question in English",
      "options_ar": ["الخيار 1","الخيار 2","الخيار 3","الخيار 4"],
      "options_en": ["Option 1","Option 2","Option 3","Option 4"],
      "answer": 0,
      "explanation_ar": "شرح الإجابة بالعربية",
      "explanation_en": "Explanation in English"
    }}
  ]
}}

Rules:
1. Exactly 10 questions.
2. Each question must be answerable from the material.
3. Do not invent facts.
4. Use different important concepts.
5. Four options for MCQ.
6. answer is zero-based correct option index.
7. Include Arabic and English for every question and option.
8. Include a short Arabic and English explanation.
"""
            raw = process_uploaded_file(uploaded, prompt)
            exam_data = extract_json(raw)

            if len(exam_data.get("questions", [])) != 10:
                raise ValueError("AI did not generate exactly 10 questions.")

            for q in exam_data["questions"]:
                q.setdefault("question_ar", q.get("question", ""))
                q.setdefault("question_en", q.get("question_ar", ""))
                q.setdefault("options_ar", q.get("options", []))
                q.setdefault("options_en", q.get("options", []))
                q.setdefault("explanation_ar", q.get("explanation", ""))
                q.setdefault("explanation_en", q.get("explanation", ""))

            session["current_exam"] = exam_data
            item_id = save_generated_item(
                session["student_id"], "exam",
                exam_data.get("title", "AI Exam"),
                json.dumps(exam_data, ensure_ascii=False)
            )
        except Exception as error:
            flash("حدث خطأ في إنشاء الاختبار: " + str(error), "error")
            return redirect(url_for("exam"))

    body = """
<div class="hero">
<h1>📝 الاختبارات الذكية</h1>
<p>10 أسئلة جامعية — السؤال والخيارات بالعربي والإنكليزي — مع التصحيح والشرح.</p>
</div>
<div class="card">
<form method="POST" enctype="multipart/form-data">
<div class="upload-box">
<h3>📚 رفع المادة</h3>
<input type="file" name="file" accept=".pdf,.png,.jpg,.jpeg,.webp,.docx,.pptx,.xlsx,.txt,.md,.csv" required>
</div>
<button class="btn" type="submit">🚀 إنشاء 10 أسئلة</button>
</form>
</div>

{% if exam_data %}
<div class="card" style="margin-top:25px;" id="examCard">
<h2>{{ exam_data.title }}</h2>
<form id="examForm">
{% for question in exam_data.questions %}
<div class="exam-question">
<h3>{{ question.number }}. {{ question.question_ar }}</h3>
<p><strong>English:</strong> {{ question.question_en }}</p>

{% for option in question.options_ar %}
<label class="option">
<input type="radio" name="q{{ loop.index0 }}" value="{{ loop.index0 }}" required>
<strong>{{ loop.index }}.</strong> {{ option }}
{% if question.options_en|length > loop.index0 %}
<br><span class="small">{{ question.options_en[loop.index0] }}</span>
{% endif %}
</label>
{% endfor %}

<div class="explanation" style="display:none;margin-top:12px;">
<strong>الإجابة الصحيحة:</strong>
<span class="correct-answer">
{{ question.options_ar[question.answer] if question.answer < question.options_ar|length else '' }}
</span>
<br>{{ question.explanation_ar }}
<br><span class="small">{{ question.explanation_en }}</span>
</div>
</div>
{% endfor %}
<button class="btn" type="button" onclick="gradeExam()">✅ تصحيح الاختبار</button>
</form>
<div id="examResult" class="result" style="display:none;"></div>
</div>
{{ actions|safe }}

<script>
const examAnswers = {{ exam_data.questions | map(attribute='answer') | list | tojson }};

function gradeExam() {
    let score = 0;
    for (let i = 0; i < examAnswers.length; i++) {
        const selected = document.querySelector('input[name="q' + i + '"]:checked');
        if (selected && Number(selected.value) === Number(examAnswers[i])) score++;
    }

    document.querySelectorAll(".explanation").forEach(el => el.style.display = "block");

    const result = document.getElementById("examResult");
    result.style.display = "block";
    const percentage = score * 10;
    result.innerHTML =
        "<h2>نتيجة الاختبار</h2>" +
        "<strong>" + score + " / 10</strong><br>" +
        "<strong>" + percentage + "%</strong><br><br>" +
        (score >= 5 ? "أحسنت، واصل المراجعة والتدريب." : "راجع المادة وحاول مرة أخرى.");
    result.scrollIntoView({behavior:"smooth"});
}
</script>
{% endif %}
"""
    actions = result_actions(item_id, "examCard", "الاختبار الجامعي")
    return render_page("الاختبارات", body, exam_data=exam_data, actions=actions)


# ============================================================
# ANNOUNCEMENTS
# ============================================================

@app.route("/announcements")
@login_required
def announcements():
    db = get_db()
    announcements = db.execute(
        "SELECT * FROM announcements ORDER BY id DESC"
    ).fetchall()
    db.close()

    body = """
<div class="hero">
<h1>📢 إعلانات جامعة الأنبار</h1>
<p>آخر الأخبار والتنبيهات والإعلانات.</p>
</div>
<div class="grid">
{% for item in announcements %}
<div class="card">
<h2>{{ item.title }}</h2>
<p style="line-height:2;">{{ item.content }}</p>
<small>{{ item.created_at }}</small>
</div>
{% else %}
<div class="card"><h3>لا توجد إعلانات حالياً.</h3></div>
{% endfor %}
</div>
"""
    return render_page("الإعلانات", body, announcements=announcements)


# ============================================================
# DOWNLOAD GENERATED RESULT
# ============================================================

@app.route("/download/<int:item_id>")
@login_required
def download_item(item_id):
    db = get_db()
    item = db.execute(
        "SELECT * FROM generated_items WHERE id = ? AND student_id = ?",
        (item_id, session["student_id"])
    ).fetchone()
    db.close()

    if not item:
        flash("الملف غير موجود.", "error")
        return redirect(url_for("home"))

    return export_text(item["title"], item["content"])


# ============================================================
# ADMIN LOGIN
# ============================================================

@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        password = request.form.get("password", "")
        if password == ADMIN_PASSWORD:
            session["admin_logged_in"] = True
            return redirect(url_for("admin"))
        flash("كلمة المرور غير صحيحة.", "error")

    body = """
<div class="card form-card">
<h2>🔐 لوحة الإدارة</h2>
<form method="POST">
<div class="form-group">
<label>كلمة مرور الإدارة</label>
<input type="password" name="password" required>
</div>
<button class="btn" type="submit">دخول الإدارة</button>
</form>
</div>
"""
    return render_page("Admin Login", body)


# ============================================================
# ADMIN DASHBOARD - STATS + ATTENDANCE + SETTINGS
# ============================================================

@app.route("/admin")
@admin_required
def admin():
    db = get_db()

    students = db.execute(
        "SELECT * FROM students ORDER BY id DESC"
    ).fetchall()

    announcements = db.execute(
        "SELECT * FROM announcements ORDER BY id DESC"
    ).fetchall()

    generated_count = db.execute(
        "SELECT COUNT(*) AS total FROM generated_items"
    ).fetchone()["total"]

    total_students = db.execute(
        "SELECT COUNT(*) AS total FROM students"
    ).fetchone()["total"]

    morning_count = db.execute(
        "SELECT COUNT(*) AS total FROM students WHERE study_type='morning'"
    ).fetchone()["total"]

    evening_count = db.execute(
        "SELECT COUNT(*) AS total FROM students WHERE study_type='evening'"
    ).fetchone()["total"]

    today = datetime.now().strftime("%Y-%m-%d")
    today_count = db.execute(
        "SELECT COUNT(*) AS total FROM students WHERE substr(created_at,1,10)=?",
        (today,)
    ).fetchone()["total"]

    stage_rows = db.execute("""
        SELECT stage, COUNT(*) AS total
        FROM students
        GROUP BY stage
        ORDER BY total DESC
    """).fetchall()

    db.close()

    body = """
<div class="hero">
<h1>⚙️ لوحة التحكم</h1>
<p>إحصائيات الطلاب، قائمة الحضور، الإعلانات وإعدادات الموقع.</p>
</div>

<div class="kpi">
<div class="card stat"><strong>{{ total_students }}</strong>إجمالي الطلاب</div>
<div class="card stat"><strong>{{ today_count }}</strong>حضور اليوم</div>
<div class="card stat"><strong>{{ morning_count }}</strong>دراسة صباحية</div>
<div class="card stat"><strong>{{ evening_count }}</strong>دراسة مسائية</div>
<div class="card stat"><strong>{{ generated_count }}</strong>عمليات AI</div>
<div class="card stat"><strong>{{ announcements|length }}</strong>الإعلانات</div>
</div>

<div class="card">
<h2>📊 توزيع الطلاب حسب المرحلة</h2>
<div class="grid">
{% for row in stage_rows %}
<div class="card stat">
<strong>{{ row.total }}</strong>
{{ row.stage }}
</div>
{% else %}
<p>لا توجد بيانات.</p>
{% endfor %}
</div>
</div>

<div class="card" style="margin-top:25px;">
<h2>👨‍🎓 قائمة الحضور</h2>
<div style="overflow:auto;">
<table class="admin-table">
<thead><tr>
<th>#</th><th>اسم الطالب</th><th>المرحلة</th>
<th>الدراسة</th><th>وقت الحضور</th>
</tr></thead>
<tbody>
{% for student in students %}
<tr>
<td>{{ student.id }}</td>
<td>{{ student.full_name }}</td>
<td>{{ student.stage }}</td>
<td>{{ "صباحي" if student.study_type == "morning" else "مسائي" }}</td>
<td>{{ student.created_at }}</td>
</tr>
{% else %}
<tr><td colspan="5">لا يوجد حضور مسجل.</td></tr>
{% endfor %}
</tbody>
</table>
</div>
</div>

<div class="card" style="margin-top:25px;">
<h2>⚙️ إعدادات الموقع</h2>
<form method="POST" action="{{ url_for('update_settings') }}" enctype="multipart/form-data">

<div class="form-group">
<label>اسم الموقع</label>
<input name="site_name" value="{{ settings.site_name }}" required>
</div>

<div class="form-group">
<label>رابط الشعار (اختياري)</label>
<input name="logo_url" value="{{ settings.logo_url }}" placeholder="https://...">
</div>

<div class="form-group">
<label>أو ارفع شعاراً من جهازك</label>
<input type="file" name="logo_file" accept=".png,.jpg,.jpeg,.webp">
</div>

{% if settings.logo_url %}
<img src="{{ settings.logo_url }}" alt="logo" style="width:90px;height:90px;object-fit:cover;border-radius:50%;margin:8px 0;">
{% endif %}

<div class="form-group">
<label>لون الموقع الرئيسي</label>
<input type="color" name="primary_color" value="{{ settings.primary_color }}">
</div>

<div class="form-group">
<label>لون الخلفية/الهيدر</label>
<input type="color" name="secondary_color" value="{{ settings.secondary_color }}">
</div>

<div class="form-group">
<label>إشعار الجامعة</label>
<input name="university_notice" value="{{ settings.university_notice }}">
</div>

<button class="btn" type="submit">💾 حفظ الإعدادات</button>
</form>
</div>

<div class="card" style="margin-top:25px;">
<h2>📢 إضافة إعلان</h2>
<form method="POST" action="{{ url_for('add_announcement') }}">
<div class="form-group"><label>عنوان الإعلان</label><input name="title" required></div>
<div class="form-group"><label>نص الإعلان</label><textarea name="content" required></textarea></div>
<button class="btn" type="submit">📢 نشر الإعلان</button>
</form>
</div>

<div style="margin-top:20px;">
<a class="btn" href="{{ url_for('admin_logout') }}">تسجيل خروج الإدارة</a>
</div>
"""
    return render_page(
        "لوحة التحكم", body,
        students=students,
        announcements=announcements,
        generated_count=generated_count,
        total_students=total_students,
        morning_count=morning_count,
        evening_count=evening_count,
        today_count=today_count,
        stage_rows=stage_rows
    )


@app.route("/admin/settings", methods=["POST"])
@admin_required
def update_settings():
    site_name = request.form.get("site_name", "").strip()
    logo_url = request.form.get("logo_url", "").strip()
    primary_color = request.form.get("primary_color", "#0d6efd").strip()
    secondary_color = request.form.get("secondary_color", "#07111f").strip()
    university_notice = request.form.get(
        "university_notice", "جامعة الأنبار"
    ).strip()

    logo_file = request.files.get("logo_file")
    if logo_file and logo_file.filename:
        filename = secure_filename(logo_file.filename)
        extension = filename.rsplit(".", 1)[1].lower() if "." in filename else ""
        if extension not in {"png", "jpg", "jpeg", "webp"}:
            flash("صيغة الشعار غير مسموحة. استخدم PNG أو JPG أو WEBP.", "error")
            return redirect(url_for("admin"))

        saved_name = (
            "site_logo_"
            + datetime.now().strftime("%Y%m%d%H%M%S")
            + "_"
            + filename
        )
        logo_file.save(os.path.join(UPLOAD_FOLDER, saved_name))
        logo_url = url_for("uploaded_file", filename=saved_name)

    db = get_db()
    db.execute("""
        UPDATE settings
        SET site_name=?, logo_url=?, primary_color=?,
            secondary_color=?, university_notice=?
        WHERE id=1
    """, (
        site_name, logo_url, primary_color,
        secondary_color, university_notice
    ))
    db.commit()
    db.close()

    flash("تم تحديث اسم الموقع والشعار والألوان بنجاح.", "success")
    return redirect(url_for("admin"))


@app.route("/admin/announcement", methods=["POST"])
@admin_required
def add_announcement():
    title = request.form.get("title", "").strip()
    content = request.form.get("content", "").strip()

    if not title or not content:
        flash("يرجى كتابة عنوان الإعلان ومحتواه.", "error")
        return redirect(url_for("admin"))

    db = get_db()
    db.execute("""
        INSERT INTO announcements (title, content, created_at)
        VALUES (?, ?, ?)
    """, (
        title, content,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ))
    db.commit()
    db.close()

    flash("تم نشر الإعلان.", "success")
    return redirect(url_for("admin"))


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("loading"))


@app.route("/admin/logout")
def admin_logout():
    session.pop("admin_logged_in", None)
    return redirect(url_for("admin_login"))


# ============================================================
# ERRORS
# ============================================================

@app.errorhandler(413)
def file_too_large(error):
    return render_page(
        "File Too Large",
        """
<div class="card">
<h2>حجم الملف كبير</h2>
<p>الحد الأقصى للملف هو 20 MB.</p>
<a class="btn" href="javascript:history.back()">رجوع</a>
</div>
"""
    ), 413


@app.errorhandler(404)
def not_found(error):
    return render_page(
        "404",
        """
<div class="card">
<h2>الصفحة غير موجودة</h2>
<a class="btn" href="/">العودة للرئيسية</a>
</div>
"""
    ), 404


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=False)
